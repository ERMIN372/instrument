"""Сквозные тесты API на настоящем PostgreSQL.

Запуск: TEST_DATABASE_URL=postgresql://user@host/db pytest
БД очищается перед каждым тестом. Без переменной тесты пропускаются.
"""
import datetime as dt
import io
import os

import pytest
from openpyxl import Workbook, load_workbook

DB_URL = os.environ.get("TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not DB_URL, reason="нужна TEST_DATABASE_URL")

HEADER = ["Дата", "Код номенклатуры", "Наименование номенклатуры", "Единица измерения",
          "Количество", "Количество базовых"]
BUN = ("001", "Булка с маком 80 г, упак 20 шт")
LOAF = ("002", "Хлеб злаковый 280 г, упак 4 шт")


def xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for day, (code, name), qty in rows:
        ws.append([day.strftime("%d.%m.%Y"), code, name, "шт", qty, qty])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def days(start: dt.date, n: int):
    return [start + dt.timedelta(days=i) for i in range(n)]


W38 = dt.date(2026, 9, 14)
W39 = dt.date(2026, 9, 21)


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setenv("DATABASE_URL", DB_URL)
    monkeypatch.setenv("APP_PASSWORD", "pw")
    import psycopg

    with psycopg.connect(DB_URL, autocommit=True) as conn:
        conn.execute("DROP TABLE IF EXISTS movements, uploads, items, sources, settings CASCADE")

    import importlib

    from app import db
    importlib.reload(db)  # пул создаётся на импорте с текущим DATABASE_URL
    from app import main
    importlib.reload(main)
    from fastapi.testclient import TestClient

    with TestClient(main.app) as c:
        c.auth = ("admin", "pw")
        yield c


def upload(c, source, content, name="file.xlsx"):
    r = c.post("/api/upload", files=[("files", (name, content))], data={"sources": [source]})
    res = r.json()["results"][0]
    assert res["ok"], res
    return res


def test_auth_required(client):
    client.auth = None
    assert client.get("/api/meta").status_code == 401
    assert client.get("/healthz").status_code == 200


def test_pivot_week_day_and_replace(client):
    # Выпуск: булка 10/день две недели, хлеб только в W39.
    upload(client, "Выпуск", xlsx(
        [(d, BUN, 10) for d in days(W38, 14)] + [(d, LOAF, 5) for d in days(W39, 7)]))
    # Остатки: срез на каждый день 14.09–28.09, суммировать нельзя.
    upload(client, "Остатки на начало", xlsx([(d, BUN, 100 + i) for i, d in enumerate(days(W38, 15))]))

    meta = client.get("/api/meta").json()
    assert [s["name"] for s in meta["sources"]] == ["Выпуск", "Остатки на начало"]
    assert [s["agg"] for s in meta["sources"]] == ["sum", "last"]  # «остат» в имени
    assert meta["sources"][1]["close_label"] == "Остатки на конец"
    assert [w["iso"] for w in meta["weeks"]] == ["2026-W40", "2026-W39", "2026-W38"]

    # Неделя 21–27.09: на начало — срез 21.09, на конец — 28.09; колонка «на конец» — последняя.
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-23"}).json()
    assert [(c["name"], c["kind"], c["date"]) for c in p["columns"]] == [
        ("Выпуск", "sum", None), ("Остатки на начало", "open", "2026-09-21"),
        ("Остатки на конец", "close", "2026-09-28")]
    rows = {r["code"]: r for r in p["rows"]}
    bun = rows["001"]["values"]
    assert bun[0] == {"cur": 70, "prev": 70, "delta": 0}
    assert bun[1]["cur"] == 107 and bun[1]["prev"] == 100  # 21.09 и 14.09
    assert bun[2]["cur"] == 114 and bun[2]["prev"] == 107  # 28.09 и 21.09
    assert rows["002"]["values"][0]["cur"] == 35 and rows["002"]["values"][0]["delta"] is None
    assert [(c["covered"], c["of"]) for c in p["columns"]] == [(7, 7), (1, 1), (1, 1)]

    # Среза на 05.10 нет — «на конец» пустая и помечена.
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-28"}).json()
    assert [(c["covered"], c["of"]) for c in p["columns"]] == [(0, 7), (1, 1), (0, 1)]
    assert {r["code"]: r for r in p["rows"]}["001"]["values"][2]["cur"] is None

    d = client.get("/api/pivot", params={"mode": "day", "date": "2026-09-23"}).json()
    bun = {r["code"]: r for r in d["rows"]}["001"]["values"]
    assert bun[0]["cur"] == 10 and bun[1]["cur"] == 109 and bun[1]["prev"] == 102
    assert bun[2]["cur"] == 110 and bun[2]["prev"] == 103  # на начало следующего дня

    # Повторная загрузка W39 с другими цифрами заменяет только W39.
    res = upload(client, "Выпуск", xlsx([(d, BUN, 1) for d in days(W39, 7)]))
    assert res["replaced"] == 14
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-21"}).json()
    rows = {r["code"]: r for r in p["rows"]}
    assert rows["001"]["values"][0]["cur"] == 7
    assert rows["001"]["values"][0]["prev"] == 70
    assert "002" not in rows  # хлеба больше нет ни в W39, ни в W38 — строка не нужна


def test_pivot_range(client):
    # Выпуск 10/день 14–27.09, остатки — срез на каждый день 14–28.09 (100, 101, …).
    upload(client, "Выпуск", xlsx([(d, BUN, 10) for d in days(W38, 14)]))
    upload(client, "Остатки на начало", xlsx([(d, BUN, 100 + i) for i, d in enumerate(days(W38, 15))]))

    # 16–20.09 сравниваем с 11–15.09 (столько же дней перед периодом); выпуск там только 14 и 15.09.
    p = client.get("/api/pivot", params={"mode": "range", "date": "2026-09-16", "date_to": "2026-09-20"}).json()
    assert p["period"]["label"] == "16.09.2026–20.09.2026 (5 дн.)"
    assert p["period"]["compare"] == "к пред. 5 дн. (11.09–15.09)"
    assert [(c["kind"], c["date"], c["covered"], c["of"]) for c in p["columns"]] == [
        ("sum", None, 5, 5), ("open", "2026-09-16", 1, 1), ("close", "2026-09-21", 1, 1)]
    bun = p["rows"][0]["values"]
    assert bun[0] == {"cur": 50, "prev": 20, "delta": 1.5}
    assert (bun[1]["cur"], bun[1]["prev"]) == (102, None)  # 16.09; среза на 11.09 нет
    assert (bun[2]["cur"], bun[2]["prev"]) == (107, 102)   # 21.09 и 16.09

    # Период ровно в неделю Пн–Вс считается как неделя.
    r = client.get("/api/pivot", params={"mode": "range", "date": "2026-09-21", "date_to": "2026-09-27"}).json()
    w = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-21"}).json()
    assert r["rows"] == w["rows"]

    sid = p["columns"][0]["id"]
    b = client.get("/api/by-days", params={"source_id": sid, "date": "2026-09-16", "date_to": "2026-09-20"}).json()
    assert b["rows"][0]["days"] == [10] * 5 and b["rows"][0]["total"] == 50 and b["rows"][0]["prev"] == 20

    x = client.get("/api/export.xlsx", params={"mode": "range", "date": "2026-09-16", "date_to": "2026-09-20"})
    assert x.status_code == 200
    assert "instrument_2026-09-16_2026-09-20.xlsx" in x.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(x.content))["По дням · Выпуск"]
    assert [ws.cell(4, c).value for c in range(9, 13)] == ["Вс 20.09", "Итого", "Пред. период", "Δ"]
    assert ws["J5"].value == '=IF(COUNT(E5:I5)=0,"",SUM(E5:I5))'
    assert ws["L5"].value.startswith('=IF(OR(K5=""')

    bad = [{"mode": "range", "date": "2026-09-16"},
           {"mode": "range", "date": "2026-09-20", "date_to": "2026-09-16"},
           {"mode": "range", "date": "2025-01-01", "date_to": "2026-09-16"},
           {"mode": "month", "date": "2026-09-16"}]
    for params in bad:
        assert client.get("/api/pivot", params=params).status_code == 400, params
        assert client.get("/api/export.xlsx", params=params).status_code == 400, params


def test_by_days_trend_item_and_export(client):
    upload(client, "Выпуск", xlsx([(d, BUN, i) for i, d in enumerate(days(W38, 14))]))
    sid = client.get("/api/meta").json()["sources"][0]["id"]

    b = client.get("/api/by-days", params={"source_id": sid, "date": "2026-09-21"}).json()
    row = b["rows"][0]
    assert row["days"] == [7, 8, 9, 10, 11, 12, 13] and row["total"] == 70 and row["prev"] == 21

    t = client.get("/api/trend", params={"source_id": sid, "end": "2026-09-21", "count": 3}).json()
    assert t["rows"][0]["values"] == [None, 21, 70]

    it = client.get("/api/item/001", params={"date": "2026-09-21"}).json()
    assert it["sources"][0]["total"] == 70 and it["sources"][0]["weeks"][-2]["value"] == 21

    x = client.get("/api/export.xlsx", params={"mode": "week", "date": "2026-09-21"})
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Сводная", "По дням · Выпуск"]
    ws = wb["Сводная"]
    assert [ws.cell(4, c).value for c in range(1, 8)] == [
        "Категория", "Код", "Номенклатура", "Ед.", "Выпуск", "Выпуск · пред.", "Выпуск · Δ"]
    assert (ws["E5"].value, ws["F5"].value) == (70, 21)
    assert ws["G5"].value.startswith("=IF(")
    assert ws["E7"].value == "=SUMIFS(E5:E5,$D$5:$D$5,$D7)"


def test_errors(client):
    assert client.get("/api/pivot", params={"date": "nope"}).status_code == 400
    assert client.get("/api/by-days", params={"source_id": 42}).status_code == 404
    r = client.post("/api/upload", files=[("files", ("bad.xlsx", b"not an xlsx"))])
    assert r.json()["results"][0]["ok"] is False
    upload(client, "А", xlsx([(W39, BUN, 1)]))
    upload(client, "Б", xlsx([(W39, BUN, 1)]))
    ids = [s["id"] for s in client.get("/api/meta").json()["sources"]]
    assert client.patch(f"/api/sources/{ids[1]}", json={"name": "А"}).status_code == 409
    assert client.patch(f"/api/sources/{ids[1]}", json={"agg": "avg"}).status_code == 400
    assert client.patch(f"/api/sources/{ids[1]}", json={"hidden": True}).status_code == 200
    p = client.get("/api/pivot", params={"date": "2026-09-21"}).json()
    assert [c["name"] for c in p["columns"]] == ["А"]


def test_rename_keeps_old_name_as_alias(client):
    upload(client, "ЗаказыПоДням", xlsx([(W39, BUN, 5)]))
    sid = client.get("/api/meta").json()["sources"][0]["id"]
    assert client.patch(f"/api/sources/{sid}", json={"name": "Заказ покупателей"}).status_code == 200
    # Файл со старым именем попадает в переименованную колонку, а не создаёт новую.
    assert client.get("/api/source-name", params={"filename": "ЗаказыПоДням_сентябрь_2026.xlsx"}).json() == {
        "source": "Заказ покупателей"}
    res = upload(client, "", xlsx([(W39 + dt.timedelta(days=1), BUN, 7)]), name="заказыподнЯм_2026.xlsx")
    assert res["source"] == "Заказ покупателей"
    meta = client.get("/api/meta").json()
    assert [(s["name"], s["aliases"]) for s in meta["sources"]] == [("Заказ покупателей", ["ЗаказыПоДням"])]
    # Возврат старого имени убирает его из алиасов, новое уходит в алиасы.
    client.patch(f"/api/sources/{sid}", json={"name": "ЗаказыПоДням"})
    assert client.get("/api/meta").json()["sources"][0]["aliases"] == ["Заказ покупателей"]


def test_rename_migration(client):
    for name in ["ЗаказыПоДням", "ОстаткиПоДням", "ПланПроизводстваПоДням", "Прочее",
                 "ФактПроизводстваПоДням", "ПланированиеПроизводства"]:
        upload(client, name, xlsx([(W39, BUN, 1)]))
    from app import db

    with db.pool.connection() as conn:
        db.rename_sources(conn)
    meta = client.get("/api/meta").json()
    assert [s["name"] for s in meta["sources"]] == [
        "Остатки на начало периода", "Заказ склада", "План производства",
        "Выпуск производства", "Заказ покупателей", "Прочее"]
    assert meta["sources"][0]["close_label"] == "Остатки на конец периода"
    p = client.get("/api/pivot", params={"date": "2026-09-21"}).json()
    assert [c["name"] for c in p["columns"]][-3:] == [  # за остатками на конец — расход
        "Остатки на конец периода", "Расход нед. 37", "Расход нед. 38"]

    # Повторный запуск ничего не трогает: ручной порядок сохраняется.
    ids = [s["id"] for s in meta["sources"]]
    client.patch(f"/api/sources/{ids[5]}", json={"position": 0})
    with db.pool.connection() as conn:
        db.rename_sources(conn)
    assert client.get("/api/meta").json()["sources"][0]["name"] == "Прочее"


def test_frontend_revalidated_after_deploy(client):
    r = client.get("/static/app.js")
    assert r.status_code == 200 and r.headers["cache-control"] == "no-cache"
    assert client.get("/static/app.js", headers={"if-none-match": r.headers["etag"]}).status_code == 304
    assert client.get("/").headers["cache-control"] == "no-cache"
    assert "cache-control" not in client.get("/api/meta").headers


def test_end_of_period(client):
    # Заказ склада: булка 10..24 по дням 14.09–28.09, хлеб — только 21.09 (на конец недели его нет).
    upload(client, "Заказ склада", xlsx(
        [(d, BUN, 10 + i) for i, d in enumerate(days(W38, 15))] + [(W39, LOAF, 3)]))
    from app import db

    # Старая БД: ограничение без end — миграция расширяет его и переводит «Заказ склада» на end.
    with db.pool.connection() as conn:
        conn.execute("ALTER TABLE sources DROP CONSTRAINT sources_agg_check")
        conn.execute("ALTER TABLE sources ADD CONSTRAINT sources_agg_check CHECK (agg IN ('sum', 'last'))")
        conn.commit()
        db.add_end_agg(conn)
    sid = client.get("/api/meta").json()["sources"][0]["id"]
    assert client.get("/api/meta").json()["sources"][0]["agg"] == "end"

    # Неделя 21–27.09: на текущий период — срез на 21.09 (пред. — 14.09), на будущий — на 28.09
    # (пред. — 21.09).
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-23"}).json()
    assert [(c["kind"], c["date"], c["covered"]) for c in p["columns"]][:2] == [
        ("current", "2026-09-21", 1), ("end", "2026-09-28", 1)]
    assert [c["name"] for c in p["columns"]][:2] == ["Заказ склада на текущий период", "Заказ склада на будущий период"]
    rows = {r["code"]: r for r in p["rows"]}
    assert rows["001"]["values"][0] == {"cur": 17, "prev": 10, "delta": (17 - 10) / 10}
    assert rows["001"]["values"][1] == {"cur": 24, "prev": 17, "delta": (24 - 17) / 17}
    assert rows["002"]["values"][:2] == [{"cur": 3, "prev": None, "delta": None}, {"cur": None, "prev": 3, "delta": None}]

    # День 23.09 → срез на 24.09; неделя 28.09 → среза на 05.10 нет.
    d = client.get("/api/pivot", params={"mode": "day", "date": "2026-09-23"}).json()
    assert [c["date"] for c in d["columns"]][:2] == ["2026-09-23", "2026-09-24"]
    assert [v["cur"] for v in d["rows"][0]["values"]][:2] == [19, 20]
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-28"}).json()
    assert (p["columns"][1]["date"], p["columns"][1]["covered"]) == ("2026-10-05", 0)

    b = client.get("/api/by-days", params={"source_id": sid, "date": "2026-09-21"}).json()
    assert {r["code"]: (r["total"], r["prev"]) for r in b["rows"]} == {"001": (24, 17), "002": (None, 3)}
    t = client.get("/api/trend", params={"source_id": sid, "end": "2026-09-21", "count": 2}).json()
    assert {r["code"]: r["values"] for r in t["rows"]}["001"] == [17, 24]
    it = client.get("/api/item/001", params={"date": "2026-09-21"}).json()
    assert [x["total"] for x in it["sources"]] == [17, 24]  # на текущий период (пн 21.09), на будущий (пн 28.09)
    x = client.get("/api/export.xlsx", params={"mode": "week", "date": "2026-09-21"})
    ws = load_workbook(io.BytesIO(x.content))["По дням · Заказ склада"]
    assert "на будущий период" in ws["A2"].value and ws["L5"].value == 24

    # Ручной выбор не перетирается повторным запуском миграции.
    assert client.patch(f"/api/sources/{sid}", json={"agg": "sum"}).status_code == 200
    with db.pool.connection() as conn:
        db.add_end_agg(conn)
    assert client.get("/api/meta").json()["sources"][0]["agg"] == "sum"


def test_mail_attachment_lands_in_pivot(client):
    from email.message import EmailMessage

    from app import mail_import

    msg = EmailMessage()
    msg["From"] = "robot@firma.ru"
    msg.set_content("выгрузка")
    msg.add_attachment(xlsx([(d, BUN, 3) for d in days(W39, 7)]), maintype="application",
                       subtype="octet-stream", filename="Выпуск_сентябрь_2026.xlsx")
    res = mail_import.process_message(msg.as_bytes(), ["robot@firma.ru"])
    assert res[0]["ok"] and res[0]["source"] == "Выпуск"

    p = client.get("/api/pivot", params={"mode": "week", "date": W39.isoformat()}).json()
    assert [c["name"] for c in p["columns"]] == ["Выпуск"]
    mail = client.get("/api/mail").json()
    assert mail["enabled"] is False
    ev = mail["events"][0]
    assert ev["kind"] == "file" and ev["ok"] and ev["source"] == "Выпуск" and ev["file"] == "Выпуск_сентябрь_2026.xlsx"

    upload(client, "Выпуск", xlsx([(d, BUN, 4) for d in days(W38, 7)]), name="руками.xlsx")
    via = {u["filename"]: u["via"] for u in client.get("/api/uploads").json()}
    assert via == {"Выпуск_сентябрь_2026.xlsx": "mail", "руками.xlsx": "site"}


def test_rc_tab_wed_to_mon(client):
    # Остаток: на каждый день 21.09–28.09 значение = число месяца; выпуск 5/день.
    # Заказ покупателей 2/день с 05.09: неделя 36 — 2 дня из 7, недели 37–38 полные.
    upload(client, "Остатки на начало периода", xlsx([(d, BUN, d.day) for d in days(W39, 8)]))
    upload(client, "Заказ покупателей", xlsx([(d, BUN, 2) for d in days(dt.date(2026, 9, 5), 23)]))
    upload(client, "Заказ склада", xlsx([(d, BUN, 100) for d in days(W39, 7)]))
    upload(client, "Выпуск производства", xlsx([(d, BUN, 5) for d in days(W39, 7)]))

    t = client.get("/api/rc", params={"date": "2026-09-25"}).json()
    assert {k: v["name"] for k, v in t["roles"].items()} == {
        "stock": "Остатки на начало периода", "order": "Заказ покупателей",
        "output": "Выпуск производства", "consumption": "Заказ покупателей"}
    cols = [(c["role"], c["sub"]) for c in t["columns"]]
    assert cols == [("stock", "Ср 23.09")] \
        + [("order", f"{wd} {d}.09") for wd, d in zip(("Ср", "Чт", "Пт", "Сб", "Вс"), (23, 24, 25, 26, 27))] \
        + [("output", f"{wd} {d}.09") for wd, d in zip(("Ср", "Чт", "Пт", "Сб", "Вс"), (23, 24, 25, 26, 27))] \
        + [("calc", "Пн 28.09"), ("consumption", "31.08–06.09"), ("consumption", "07.09–13.09"),
           ("consumption", "14.09–20.09")]
    assert [c["label"] for c in t["columns"][-3:]] == ["Потребл. нед. 36", "Потребл. нед. 37", "Потребл. нед. 38"]
    assert [(c["covered"], c["of"]) for c in t["columns"][-4:]] == [(1, 1), (2, 7), (7, 7), (7, 7)]
    # Остаток пн — расчёт, а не срез 1С (28): остаток ср на начало дня, поэтому движение
    # самой среды тоже входит: 23 − 5×2 + 5×5 = 38.
    assert t["rows"][0]["values"] == [23, 2, 2, 2, 2, 2, 5, 5, 5, 5, 5, 38, 4, 14, 14]
    assert t["period"]["label"].startswith("Ср 23.09 → Пн 28.09")

    # Заказ — другой источник, выбор общий и сохраняется.
    sklad = next(s["id"] for s in client.get("/api/meta").json()["sources"] if s["name"] == "Заказ склада")
    assert client.put("/api/rc-settings", json={"order": sklad}).json()["order"] == sklad
    t = client.get("/api/rc", params={"date": "2026-09-25"}).json()
    assert t["roles"]["order"]["name"] == "Заказ склада"
    assert t["rows"][0]["values"][1:6] == [100] * 5
    assert t["rows"][0]["values"][11] == 23 - 500 + 25  # расчёт идёт от выбранного заказа
    assert t["roles"]["consumption"]["name"] == "Заказ покупателей"  # остальные роли не сбились

    # Неделя 38: остатка на среду 16.09 нет — расчёт не делаем; потребление — недели 35–37.
    t = client.get("/api/rc", params={"date": "2026-09-16"}).json()
    assert t["rows"][0]["values"] == [None] * 13 + [4, 14]
    assert not t["columns"][0]["covered"] and not t["columns"][11]["covered"]

    # Выпуск только 21 и 27.09: в загруженном периоде дни без строк — это ноль, а не «нет данных».
    upload(client, "Выпуск производства", xlsx([(W39, BUN, 5), (dt.date(2026, 9, 27), BUN, 5)]))
    t = client.get("/api/rc", params={"date": "2026-09-25"}).json()
    out = [c for c in t["columns"] if c["role"] == "output"]
    assert all(c["covered"] for c in out)
    assert t["rows"][0]["values"][6:11] == [None, None, None, None, 5]
    upload(client, "Выпуск производства", xlsx([(d, BUN, 5) for d in days(W39, 7)]))

    r = client.get("/api/export-rc.xlsx", params={"date": "2026-09-25"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Товародвиженец РЦ"
    assert ws["E4"].value == "Остаток Ср 23.09" and ws["F4"].value == "Заказ Ср 23.09"
    assert ws["P4"].value == "Остаток расчёт Пн 28.09" and ws["S4"].value == "Потребл. нед. 38 14.09–20.09"
    assert [ws.cell(row=5, column=c).value for c in range(5, 20)] == \
        [23, 100, 100, 100, 100, 100, 5, 5, 5, 5, 5, 23 - 500 + 25, 4, 14, 14]


def test_order_columns_current_and_future(client):
    """Круассан со скрина, порядок колонок как просил заказчик: остаток на начало (14.09),
    заказ склада на текущий период (14.09 → 15 960), заказ покупателей, заказ склада на
    будущий период (21.09 → 11 970, под него производят эту неделю — совпадает с планом),
    план, выпуск, остаток на конец (21.09). Карточка товара — в том же порядке."""
    D = dt.date
    # Загружаем вразнобой — порядок выставляет разовая миграция.
    upload(client, "План производства", xlsx([(D(2026, 9, 17), BUN, 11970)]))
    upload(client, "Выпуск производства", xlsx([(D(2026, 9, 18), BUN, 4560)]))
    upload(client, "Заказ покупателей", xlsx([(D(2026, 9, 14), BUN, 3690), (D(2026, 9, 15), BUN, 990)]))
    upload(client, "Остатки на начало периода", xlsx([(D(2026, 9, 14), BUN, 26670), (D(2026, 9, 21), BUN, 17040)]))
    upload(client, "Заказ склада", xlsx([(D(2026, 9, 14), BUN, 15960), (D(2026, 9, 21), BUN, 11970)]))
    srcs = {x["name"]: x["id"] for x in client.get("/api/meta").json()["sources"]}
    client.patch(f"/api/sources/{srcs['Заказ склада']}", json={"agg": "end"})
    from app import db
    with db.pool.connection() as conn:
        db.order_sources(conn)

    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-09-16"}).json()
    got = [(c["name"], c.get("date"), v["cur"]) for c, v in zip(p["columns"], p["rows"][0]["values"])][:7]
    assert got == [
        ("Остатки на начало периода", "2026-09-14", 26670),
        ("Заказ склада на текущий период", "2026-09-14", 15960),
        ("Заказ покупателей", None, 3690 + 990),
        ("Заказ склада на будущий период", "2026-09-21", 11970),
        ("План производства", None, 11970),
        ("Выпуск производства", None, 4560),
        ("Остатки на конец периода", "2026-09-21", 17040),
    ]

    card = client.get(f"/api/item/{BUN[0]}", params={"date": "2026-09-16"}).json()
    assert [(x["name"], x["label"], x["total"]) for x in card["sources"]] == [
        ("Остатки на начало периода", None, 26670),
        ("Заказ склада", "на текущий период", 15960),
        ("Заказ покупателей", None, 3690 + 990),
        ("Заказ склада", "на будущий период", 11970),
        ("План производства", None, 11970),
        ("Выпуск производства", None, 4560),
    ]
    cur, fut = (x for x in card["sources"] if x["name"] == "Заказ склада")
    assert cur["days"] == [15960] + [None] * 6 and cur["weeks"] == []
    assert fut["days"] == [None] * 7 and fut["as_of"] == "2026-09-21"

    # Ручной порядок из «Источников» миграция больше не трогает.
    client.patch(f"/api/sources/{srcs['План производства']}", json={"position": 0})
    with db.pool.connection() as conn:
        db.order_sources(conn)
    assert client.get("/api/meta").json()["sources"][0]["name"] == "План производства"


def test_pivot_spent_last_two_weeks(client):
    """Расход в конце сводной: заказ покупателей за две недели до выбранной — как во вкладке
    «Динамика», сумма за Пн–Вс. Неделя 05–11.10 → нед. 39 (21–27.09) и нед. 40 (28.09–04.10),
    Δ — к неделе раньше. Заказ склада в расход не идёт."""
    D = dt.date
    upload(client, "Остатки на начало периода", xlsx([(D(2026, 10, 5), BUN, 100)]))
    upload(client, "Заказ склада", xlsx([(D(2026, 9, 21), BUN, 500), (D(2026, 9, 28), BUN, 600)]))
    upload(client, "Заказ покупателей", xlsx([
        (D(2026, 9, 15), BUN, 4), (D(2026, 9, 18), BUN, 6),                            # нед. 38: 10
        (D(2026, 9, 21), BUN, 7), (D(2026, 9, 26), BUN, 10), (D(2026, 9, 22), LOAF, 3),  # нед. 39: 17
        (D(2026, 9, 28), BUN, 20), (D(2026, 10, 4), BUN, 4),                           # нед. 40: 24
        (D(2026, 10, 5), BUN, 30)]))                                                   # сам период
    srcs = {x["name"]: x["id"] for x in client.get("/api/meta").json()["sources"]}
    sid = srcs["Заказ покупателей"]
    client.patch(f"/api/sources/{srcs['Заказ склада']}", json={"agg": "end"})

    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-10-07"}).json()
    cols = p["columns"][-2:]
    assert [(c["id"], c["name"], c["kind"], c["date"], c["covered"], c["of"], c["week"]["start"]) for c in cols] == [
        (sid, "Расход нед. 39", "spent", None, 7, 7, "2026-09-21"),
        (sid, "Расход нед. 40", "spent", None, 7, 7, "2026-09-28"),
    ]
    assert sum(c["kind"] == "spent" for c in p["columns"]) == 2
    assert p["columns"][-3]["name"] == "Остатки на конец периода"  # новые — после остатков на конец
    rows = {r["code"]: r["values"][-2:] for r in p["rows"]}
    assert rows["001"] == [
        {"cur": 17, "prev": 10, "delta": (17 - 10) / 10},
        {"cur": 24, "prev": 17, "delta": (24 - 17) / 17},
    ]
    assert rows["002"] == [{"cur": 3, "prev": None, "delta": None}, {"cur": None, "prev": 3, "delta": None}]

    # Сверка с «Динамикой» по тем же неделям — цифры должны совпадать.
    t = client.get("/api/trend", params={"source_id": sid, "end": "2026-09-28", "count": 2}).json()
    assert [w["iso"] for w in t["weeks"]] == ["2026-W39", "2026-W40"]
    trend = {r["code"]: r["values"] for r in t["rows"]}
    assert {code: [v["cur"] for v in vals] for code, vals in rows.items()} == {"001": trend["001"], "002": trend["002"]}

    # День 30.09 — две недели до его недели (28.09): нед. 38 (файл с 15.09 — 6 из 7 дн.) и 39.
    p = client.get("/api/pivot", params={"mode": "day", "date": "2026-09-30"}).json()
    assert [(c["name"], c["covered"], c["of"]) for c in p["columns"][-2:]] == [
        ("Расход нед. 38", 6, 7), ("Расход нед. 39", 7, 7)]
    # Неделя 12.10 → нед. 41: файл кончается 05.10 — 1 из 7 дн.
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-10-12"}).json()
    assert [(c["covered"], c["of"]) for c in p["columns"][-2:]] == [(7, 7), (1, 7)]

    x = client.get("/api/export.xlsx", params={"mode": "week", "date": "2026-10-07"})
    ws = load_workbook(io.BytesIO(x.content))["Сводная"]
    assert "Расход — заказ покупателей" in ws["A2"].value
    assert [c.value for c in ws[4]][-3:] == ["Расход нед. 40", "Расход нед. 40 · пред.", "Расход нед. 40 · Δ"]

    # Без заказа покупателей колонок расхода нет — заказ склада их не даёт.
    client.patch(f"/api/sources/{sid}", json={"hidden": True})
    p = client.get("/api/pivot", params={"mode": "week", "date": "2026-10-07"}).json()
    assert not any(c["kind"] == "spent" for c in p["columns"])
