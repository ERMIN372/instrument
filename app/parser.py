"""Разбор выгрузки 1С «движение номенклатуры по дням».

Ожидаемые колонки (порядок, регистр и лишние колонки не важны):
    Дата | UID номенклатуры | Код номенклатуры | Наименование номенклатуры |
    Единица измерения | Количество | Количество базовых
Обязательны: «Дата», «Наименование» или «Код», «Количество базовых» или «Количество».

Считаем в базовых единицах (шт, кг): «Количество» бывает в упаковках, и в разных
файлах упаковки разные, складывать их нельзя. Повторы (день, товар) суммируются,
строки без даты («Итого») пропускаются.
"""
from __future__ import annotations

import datetime as dt
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import PurePath

from openpyxl import load_workbook

HEADER_SCAN_ROWS = 20

# Поле -> варианты заголовка (после нормализации: нижний регистр, ё→е, одиночные пробелы).
COLUMNS = {
    "date": ("дата", "день", "период"),
    "uid": ("uid номенклатуры", "uid", "гуид"),
    "code": ("код номенклатуры", "код", "артикул"),
    "name": ("наименование номенклатуры", "номенклатура", "наименование", "товар"),
    "unit": ("единица измерения", "ед. изм.", "ед.изм.", "ед изм", "ед."),
    "qty_base": ("количество базовых", "количество в базовых единицах", "кол-во базовых"),
    "qty": ("количество", "кол-во"),
}

DATE_STR_RE = re.compile(r"^\d{1,4}[./-]\d{1,2}[./-]\d{2,4}")
DATE_FORMATS = ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
NUM_JUNK_RE = re.compile(r"[\s  ]|руб\.?|₽", re.I)
NUM_RE = re.compile(r"-?\d+(\.\d+)?")
PACK_RE = re.compile(r"упак\w*\.?,?\s*(\d+)\s*шт", re.I)
CATEGORY_RE = re.compile(r"[A-Za-zА-Яа-яЁё-]+")
MONTH_RE = re.compile(
    r"\b(январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|октябр|ноябр|декабр)\w*", re.I
)


@dataclass
class Item:
    code: str
    name: str
    uid: str | None
    base_unit: str
    pack_size: int | None
    category: str


@dataclass
class ParsedFile:
    # (день, код товара) -> [количество в единицах файла, количество базовых]
    movements: dict[tuple[dt.date, str], list[float]] = field(default_factory=dict)
    items: dict[str, Item] = field(default_factory=dict)
    rows: int = 0
    skipped: int = 0
    merged: int = 0
    negative: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def days(self) -> list[dt.date]:
        return sorted({d for d, _ in self.movements})


def to_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        d = v.date()
    elif isinstance(v, dt.date):
        d = v
    elif isinstance(v, str) and DATE_STR_RE.match(v.strip()):
        s = v.strip().split()[0]
        d = None
        for fmt in DATE_FORMATS:
            try:
                d = dt.datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
    else:
        return None
    if d is None or not 2000 <= d.year <= 2100:
        return None
    return d


def to_number(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        s = NUM_JUNK_RE.sub("", v).replace(",", ".")
        if NUM_RE.fullmatch(s):
            return float(s)
    return None


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _norm(v) -> str:
    return text(v).lower().replace("ё", "е")


def source_from_filename(filename: str) -> str:
    """«Выпуск_сентябрь_2026.xlsx» -> «Выпуск»: без дат, месяцев и номеров,
    чтобы файлы одного источника за разные периоды попадали в одну колонку."""
    stem = PurePath(filename).stem
    s = re.sub(r"[_\-.()\[\]]+", " ", stem)  # «_» — словесный символ, мешает \b ниже
    s = MONTH_RE.sub(" ", s)
    s = re.sub(r"\d+", " ", s)
    s = re.sub(r"\b(неделя|нед|week|w|с|по)\b", " ", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).strip()
    return s or stem


def _map_columns(header: list) -> dict[str, int]:
    """Поле -> индекс колонки. Сначала точные совпадения, потом «содержит»,
    чтобы «Количество» не съело колонку «Количество базовых»."""
    cells = {i: _norm(v) for i, v in enumerate(header) if _norm(v)}
    found: dict[str, int] = {}
    for fld, aliases in COLUMNS.items():
        for i, h in cells.items():
            if h in aliases and i not in found.values():
                found[fld] = i
                break
    for fld, aliases in COLUMNS.items():
        if fld in found:
            continue
        for i, h in cells.items():
            if i not in found.values() and any(a in h for a in aliases):
                found[fld] = i
                break
    return found


def _item_meta(name: str, unit: str) -> tuple[str, int | None, str]:
    pack = PACK_RE.search(unit) or PACK_RE.search(name)
    pack_size = int(pack.group(1)) if pack else None
    u = unit.lower()
    if not u or u.startswith("упак"):
        base_unit = "шт"
    else:
        base_unit = u.rstrip(".")
    m = CATEGORY_RE.search(name)
    category = m.group(0).capitalize() if m else "Прочее"
    return base_unit, pack_size, category


def _parse_sheet(ws, out: ParsedFile) -> bool:
    rows = ws.iter_rows(values_only=True)
    cols: dict[str, int] = {}
    for _ in range(HEADER_SCAN_ROWS):
        header = next(rows, None)
        if header is None:
            return False
        cols = _map_columns(list(header))
        if "date" in cols and ("qty_base" in cols or "qty" in cols) and ("name" in cols or "code" in cols):
            break
    else:
        return False

    def get(row, fld):
        i = cols.get(fld)
        return row[i] if i is not None and i < len(row) else None

    for row in rows:
        if not any(v not in (None, "") for v in row):
            continue
        day = to_date(get(row, "date"))
        name = text(get(row, "name"))
        code = text(get(row, "code")) or name
        qb = to_number(get(row, "qty_base"))
        q = to_number(get(row, "qty"))
        if qb is None:
            qb = q
        if q is None:
            q = qb
        if day is None or not code or qb is None:
            out.skipped += 1
            continue
        out.rows += 1
        if qb < 0:
            out.negative += 1
        if code not in out.items:
            base_unit, pack_size, category = _item_meta(name or code, text(get(row, "unit")))
            out.items[code] = Item(code, name or code, text(get(row, "uid")) or None, base_unit, pack_size, category)
        key = (day, code)
        if key in out.movements:
            out.merged += 1
            out.movements[key][0] += q
            out.movements[key][1] += qb
        else:
            out.movements[key] = [q, qb]
    return True


def parse_xlsx(data: bytes) -> ParsedFile:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = ParsedFile()
    try:
        found = False
        for ws in wb.worksheets:
            if hasattr(ws, "reset_dimensions"):
                ws.reset_dimensions()  # иначе read-only режим верит битому <dimension> в файле
            found = _parse_sheet(ws, out) or found
    finally:
        wb.close()
    if not found:
        raise ValueError(
            "не нашёл строку заголовка с колонками «Дата», «Наименование/Код» и «Количество»"
        )
    if not out.movements:
        raise ValueError("заголовок есть, но нет ни одной строки с датой и количеством")
    if out.merged:
        out.warnings.append(f"{out.merged} повторов «день + товар» сложены")
    if out.negative:
        out.warnings.append(f"{out.negative} строк с отрицательным количеством")
    if out.skipped:
        out.warnings.append(f"{out.skipped} строк без даты/товара/количества пропущено")
    return out
