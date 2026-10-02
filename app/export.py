"""Выгрузка сводной и раскладки по дням в xlsx. Итоги и Δ — формулами Excel."""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="E8EAF0")
NUM = "#,##0.##;-#,##0.##;0"
PCT = "+0.0%;-0.0%;0.0%"
FIXED = ["Категория", "Код", "Номенклатура", "Ед."]


def filter_rows(rows: list[dict], category: str | None, q: str | None) -> list[dict]:
    q = (q or "").strip().lower()
    return [
        r for r in rows
        if (not category or r["category"] == category)
        and (not q or q in r["name"].lower() or q in r["code"].lower())
    ]


def _sheet_title(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Лист"
    title, n = base, 2
    while title in used:
        suffix = f" {n}"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title)
    return title


def _header(ws, row: int, labels: list[str]):
    for c, h in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fixed_cells(ws, r: int, row: dict):
    for c, v in enumerate([row["category"], row["code"], row["name"], row["unit"]], 1):
        ws.cell(row=r, column=c, value=v).font = Font(name=FONT)


def _totals(ws, first: int, last: int, units: list[str], num_cols: list[int], pct_cols: dict[int, tuple[int, int]]):
    """Строки «Итого, <ед.>» формулами SUMIFS по колонке «Ед.»; Δ — от итоговых сумм."""
    r = last + 2
    for unit in units:
        label = ws.cell(row=r, column=3, value=f"Итого, {unit}")
        label.font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=4, value=unit).font = Font(name=FONT, bold=True)
        for c in num_cols:
            col = get_column_letter(c)
            cell = ws.cell(row=r, column=c, value=f'=SUMIFS({col}{first}:{col}{last},$D${first}:$D${last},$D{r})')
            cell.font = Font(name=FONT, bold=True)
            cell.number_format = NUM
        for c, (cur_c, prev_c) in pct_cols.items():
            cur, prev = f"{get_column_letter(cur_c)}{r}", f"{get_column_letter(prev_c)}{r}"
            cell = ws.cell(row=r, column=c, value=f'=IF({prev}=0,"",({cur}-{prev})/ABS({prev}))')
            cell.font = Font(name=FONT, bold=True)
            cell.number_format = PCT
        r += 1


def _widths(ws, n_cols: int):
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 6
    for c in range(5, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14


def _pivot_sheet(ws, table: dict, rows: list[dict]):
    period = table["period"]
    ws["A1"] = f"Сводная · {period['label']} · в базовых единицах"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = (f"Δ — {period['compare']}. Остаток на начало — срез на первый день периода,"
                " на конец — на первый день следующего.")
    if any(c["kind"] == "spent" for c in table["columns"]):
        ws["A2"] = (ws["A2"].value + " Расход — заказ покупателей за две прошлые недели"
                    " (за неделю, как в «Динамике»), его Δ — к неделе раньше.")
    ws["A2"].font = Font(name=FONT, italic=True, color="6B7280", size=9)

    labels = list(FIXED)
    for s in table["columns"]:
        labels += [s["name"], f"{s['name']} · пред.", f"{s['name']} · Δ"]
    _header(ws, 4, labels)

    first = r = 5
    for row in rows:
        _fixed_cells(ws, r, row)
        for i, v in enumerate(row["values"]):
            c = 5 + i * 3
            for off, val in ((0, v["cur"]), (1, v["prev"])):
                cell = ws.cell(row=r, column=c + off, value=val)
                cell.font = Font(name=FONT, bold=(off == 0))
                cell.number_format = NUM
            cur, prev = f"{get_column_letter(c)}{r}", f"{get_column_letter(c + 1)}{r}"
            d = ws.cell(row=r, column=c + 2, value=f'=IF(OR({prev}="",{prev}=0,{cur}=""),"",({cur}-{prev})/ABS({prev}))')
            d.font = Font(name=FONT)
            d.number_format = PCT
        r += 1
    last = max(first, r - 1)

    n = len(table["columns"])
    units = sorted({row["unit"] for row in rows})
    num_cols = [5 + i * 3 + off for i in range(n) for off in (0, 1)]
    pct_cols = {5 + i * 3 + 2: (5 + i * 3, 5 + i * 3 + 1) for i in range(n)}
    _totals(ws, first, last, units, num_cols, pct_cols)
    _widths(ws, len(labels))
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(labels))}{last}"


def _days_sheet(ws, table: dict, rows: list[dict]):
    src = table["source"]
    week = "iso" in table["period"]  # иначе — произвольный период
    if week:
        how = {"last": "остаток на начало недели (срез на понедельник)",
               "end": "на будущий период (срез на понедельник следующей недели)"}.get(src["agg"], "сумма за неделю")
        total_h, prev_h, note = "Итого нед.", "Пред. нед.", ""
    else:
        how = {"last": "остаток на начало периода (срез на первый день)",
               "end": "на будущий период (срез на день после конца периода)"}.get(src["agg"], "сумма за период")
        total_h, prev_h, note = "Итого", "Пред. период", f" Δ — {table['period']['compare']}."
    ws["A1"] = f"{src['name']} по дням · {table['period']['label']}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = f"{total_h} — {how}. В базовых единицах.{note}"
    ws["A2"].font = Font(name=FONT, italic=True, color="6B7280", size=9)

    n = len(table["days"])
    tot_c, prev_c, d_c = 5 + n, 6 + n, 7 + n
    tot, prv = get_column_letter(tot_c), get_column_letter(prev_c)
    labels = FIXED + [d["short"] for d in table["days"]] + [total_h, prev_h, "Δ"]
    _header(ws, 4, labels)
    first = r = 5
    for row in rows:
        _fixed_cells(ws, r, row)
        for i, v in enumerate(row["days"]):
            ws.cell(row=r, column=5 + i, value=v).number_format = NUM
        span = f"E{r}:{get_column_letter(4 + n)}{r}"
        total = f'=IF(COUNT({span})=0,"",SUM({span}))' if src["agg"] == "sum" else row["total"]
        cell = ws.cell(row=r, column=tot_c, value=total)
        cell.font = Font(name=FONT, bold=True)
        cell.number_format = NUM
        ws.cell(row=r, column=prev_c, value=row["prev"]).number_format = NUM
        d = ws.cell(row=r, column=d_c,
                    value=f'=IF(OR({prv}{r}="",{prv}{r}=0,{tot}{r}=""),"",({tot}{r}-{prv}{r})/ABS({prv}{r}))')
        d.number_format = PCT
        for c in range(5, d_c + 1):
            if c != tot_c:
                ws.cell(row=r, column=c).font = Font(name=FONT)
        r += 1
    last = max(first, r - 1)

    units = sorted({row["unit"] for row in rows})
    _totals(ws, first, last, units, list(range(5, d_c)), {d_c: (tot_c, prev_c)})
    _widths(ws, len(labels))
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(d_c)}{last}"


def rc_workbook(table: dict, category: str | None = None, q: str | None = None) -> bytes:
    """Вкладка товародвиженца РЦ: остаток ср, заказ и выпуск Ср–Вс, остаток пн, потребление за 3 нед.
    (или то же за произвольный период — как считали, пишем во второй строке)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Товародвиженец РЦ"
    rows = filter_rows(table["rows"], category, q)
    names = {role: (src or {}).get("name", "не выбран") for role, src in table["roles"].items()}
    ws["A1"] = f"Товародвиженец РЦ · {table['period']['label']} · в базовых единицах"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = (f"{table['note']} Остаток — «{names['stock']}», заказ — «{names['order']}», выпуск — «{names['output']}», "
                f"потребление по неделям — «{names['consumption']}».")
    ws["A2"].font = Font(name=FONT, italic=True, color="6B7280", size=9)

    labels = FIXED + [f"{c['label']} {c['sub']}" for c in table["columns"]]
    _header(ws, 4, labels)
    first = r = 5
    for row in rows:
        _fixed_cells(ws, r, row)
        for i, v in enumerate(row["values"]):
            cell = ws.cell(row=r, column=5 + i, value=v)
            cell.font = Font(name=FONT)
            cell.number_format = NUM
        r += 1
    last = max(first, r - 1)
    _totals(ws, first, last, sorted({row["unit"] for row in rows}), list(range(5, len(labels) + 1)), {})
    _widths(ws, len(labels))
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(labels))}{last}"
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PCT_PLAIN = "0%"


def autoorder_workbook(table: dict, category: str | None = None, q: str | None = None) -> bytes:
    """Автозаказ (демо): по каждой формуле — заказ и «почему так»; вторым листом — прогон по истории."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Автозаказ (демо)"
    rows = filter_rows(table["rows"], category, q)
    p = table["params"]
    ws["A1"] = f"Автозаказ (демо-версия, тестируем) · {table['period']['label']} · в базовых единицах"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    formulas = table["formulas"]
    ws["A2"] = f"{table['formula']} — {table['formula_note']} {table['basis']}"
    ws["A3"] = "; ".join(f"{f['label']}: {f['rule']}" for f in formulas)
    for cell in (ws["A2"], ws["A3"]):
        cell.font = Font(name=FONT, italic=True, color="6B7280", size=9)

    labels = list(FIXED)
    for f in formulas:
        labels += [f"{f['label']} · заказ", f"{f['label']} · почему"]
    _header(ws, 4, labels)
    first = r = 5
    for row in rows:
        _fixed_cells(ws, r, row)
        for i, offer in enumerate(row["offers"]):
            cell = ws.cell(row=r, column=5 + 2 * i, value=offer["value"])
            cell.font = Font(name=FONT, bold=True)
            cell.number_format = NUM
            ws.cell(row=r, column=6 + 2 * i, value=offer["why"]).font = Font(name=FONT, color="6B7280")
        r += 1
    last = max(first, r - 1)
    orders = [5 + 2 * i for i in range(len(formulas))]
    _totals(ws, first, last, sorted({row["unit"] for row in rows}), orders, {})
    _widths(ws, len(labels))
    for c in orders:
        ws.column_dimensions[get_column_letter(c + 1)].width = 42
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(labels))}{last}"

    bt = table["backtest"]
    ws = wb.create_sheet("Прогон по истории")
    weeks = ", ".join(w["iso"].split("-")[1] for w in bt["weeks"]) or "нет прошедших недель с данными"
    ws["A1"] = f"Прогон формул по прошлым неделям: {weeks}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = (f"Покрытие = (остаток пн + заказ) / потребление этой недели, в неделях. Цель — не меньше "
                f"{p['cover_weeks']}; меньше 1 — дефицит, к концу недели остаток в минусе. "
                "Считаются товаро-недели с потреблением больше нуля.")
    ws["A2"].font = Font(name=FONT, italic=True, color="6B7280", size=9)
    _header(ws, 4, ["Формула", "Товаро-недель", "Покрытие не ниже цели", "Дефицит (покрытие < 1)",
                    "Медиана покрытия, нед."])
    for r, x in enumerate(bt["rows"], 5):
        for c, (v, f) in enumerate(((x["label"], None), (x["n"], NUM), (x["hit"], PCT_PLAIN),
                                    (x["short"], PCT_PLAIN), (x["cover"], "0.0")), 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT)
            if f:
                cell.number_format = f
    ws.column_dimensions["A"].width = 40
    for c in "BCDE":
        ws.column_dimensions[c].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def workbook(pivot: dict, by_days: list[dict], category: str | None = None, q: str | None = None) -> bytes:
    wb = Workbook()
    used: set = set()
    ws = wb.active
    ws.title = _sheet_title("Сводная", used)
    _pivot_sheet(ws, pivot, filter_rows(pivot["rows"], category, q))
    for table in by_days:
        sheet = wb.create_sheet(_sheet_title(f"По дням · {table['source']['name']}", used))
        _days_sheet(sheet, table, filter_rows(table["rows"], category, q))
    wb.calculation.fullCalcOnLoad = True  # Excel пересчитает формулы при открытии
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
